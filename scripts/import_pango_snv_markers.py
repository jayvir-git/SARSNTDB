"""Build pango designation-marker SQL from Jim Kelley's workbook.

SNVs (Jim, 2026-09-09): Column C = coordinate, D = ref, E = alt, I = ratio.
If len(ref) == 1 and len(alt) == 1 and ratio < 0.2, Column A (lineage) is
attached to that SNV.

Indels (Jim, 2026-09-09 follow-up): if len(D) > len(E) the row is a deletion;
if len(D) < len(E) it is an insertion. Same ratio < 0.2 rule. Stored in a
separate table from SNVs.

Do not parse xlsx in PHP. Re-run this script if the spreadsheet changes.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INBOX = ROOT / "_incoming" / "jim-kelley"
PACKET_FILES = INBOX / "2026-09-09_pango-snv-lineages" / "files"
SOURCE_NAME = "pango-designation-markers-v1.9-calc-sort.xlsx"
OUT_SQL = ROOT / "sql" / "snv_pango_marker.sql"
OUT_DROP_SQL = ROOT / "sql" / "snv_pango_marker_drop.sql"
OUT_INDEL_SQL = ROOT / "sql" / "pango_indel_marker.sql"
OUT_INDEL_DROP_SQL = ROOT / "sql" / "pango_indel_marker_drop.sql"

RATIO_LT = 0.2


def sql_str(value: object | None) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def sql_num(value: object | None) -> str:
    if value is None:
        return "NULL"
    return repr(float(value))


def find_xlsx() -> Path:
    named = (
        PACKET_FILES / SOURCE_NAME,
        INBOX / SOURCE_NAME,
    )
    for path in named:
        if path.is_file():
            return path
    hits = sorted(INBOX.glob("**/" + SOURCE_NAME))
    if hits:
        return hits[0]
    raise SystemExit("Missing " + SOURCE_NAME + " under " + str(INBOX))


def cell_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_ratio(value: object | None) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def is_single_base_snv(ref: str, alt: str) -> bool:
    return len(ref) == 1 and len(alt) == 1


def indel_kind(ref: str, alt: str) -> str | None:
    """Jim: len(D) > len(E) deletion; len(D) < len(E) insertion."""
    if ref == "" or alt == "":
        return None
    if len(ref) > len(alt):
        return "deletion"
    if len(ref) < len(alt):
        return "insertion"
    return None


def ratio_qualifies(ratio: float | None, threshold: float = RATIO_LT) -> bool:
    return ratio is not None and ratio < threshold


def qualifies_snv_row(ref: str, alt: str, ratio: float | None, threshold: float = RATIO_LT) -> bool:
    """True when Jim's rule attaches Column A to the SNV."""
    return is_single_base_snv(ref, alt) and ratio_qualifies(ratio, threshold)


def qualifies_indel_row(ref: str, alt: str, ratio: float | None, threshold: float = RATIO_LT) -> bool:
    """True when the row is a designation-marker indel the lineage carries."""
    return indel_kind(ref, alt) is not None and ratio_qualifies(ratio, threshold)


def _row_dict(
    lineage: str,
    chrom: str,
    coordinate: int,
    ref: str,
    alt: str,
    adref: object,
    adalt: object,
    dp: object,
    ratio: float | None,
    source_row: int,
    kind: str | None = None,
) -> dict[str, object]:
    row: dict[str, object] = {
        "lineage": lineage,
        "chrom": chrom or None,
        "coordinate": coordinate,
        "reference": ref,
        "alternate": alt,
        "adref": int(adref) if isinstance(adref, (int, float)) else None,
        "adalt": int(adalt) if isinstance(adalt, (int, float)) else None,
        "dp": int(dp) if isinstance(dp, (int, float)) else None,
        "ratio": ratio,
        "source_row": source_row,
    }
    if kind is not None:
        row["kind"] = kind
    return row


def _keep_lower_ratio(
    out: list[dict[str, object]],
    seen: dict[tuple[str, int, str, str], int],
    key: tuple[str, int, str, str],
    row: dict[str, object],
) -> None:
    prev = seen.get(key)
    ratio = row["ratio"]
    if prev is None:
        seen[key] = len(out)
        out.append(row)
        return
    if ratio is not None and (out[prev]["ratio"] is None or float(ratio) < float(out[prev]["ratio"])):
        out[prev] = row


def iter_designation_rows(path: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    snvs: list[dict[str, object]] = []
    indels: list[dict[str, object]] = []
    seen_snv: dict[tuple[str, int, str, str], int] = {}
    seen_indel: dict[tuple[str, int, str, str], int] = {}

    for source_row, raw in enumerate(ws.iter_rows(min_row=2, max_col=9, values_only=True), start=2):
        lineage = cell_text(raw[0] if len(raw) > 0 else None)
        chrom = cell_text(raw[1] if len(raw) > 1 else None)
        pos_raw = raw[2] if len(raw) > 2 else None
        ref = cell_text(raw[3] if len(raw) > 3 else None).upper()
        alt = cell_text(raw[4] if len(raw) > 4 else None).upper()
        adref = raw[5] if len(raw) > 5 else None
        adalt = raw[6] if len(raw) > 6 else None
        dp = raw[7] if len(raw) > 7 else None
        ratio = parse_ratio(raw[8] if len(raw) > 8 else None)

        if lineage == "" or pos_raw is None:
            continue
        try:
            coordinate = int(pos_raw)
        except (TypeError, ValueError):
            continue

        if qualifies_snv_row(ref, alt, ratio):
            _keep_lower_ratio(
                snvs,
                seen_snv,
                (lineage, coordinate, ref, alt),
                _row_dict(lineage, chrom, coordinate, ref, alt, adref, adalt, dp, ratio, source_row),
            )
            continue

        kind = indel_kind(ref, alt)
        if kind is not None and ratio_qualifies(ratio):
            _keep_lower_ratio(
                indels,
                seen_indel,
                (lineage, coordinate, ref, alt),
                _row_dict(
                    lineage, chrom, coordinate, ref, alt, adref, adalt, dp, ratio, source_row, kind
                ),
            )

    return snvs, indels


def iter_qualifying_rows(path: Path) -> list[dict[str, object]]:
    snvs, _indels = iter_designation_rows(path)
    return snvs


def write_drop_sql() -> None:
    OUT_DROP_SQL.write_text(
        "-- Reverse sql/snv_pango_marker.sql without touching the mutations table.\n"
        "SET FOREIGN_KEY_CHECKS=0;\n"
        "DROP TABLE IF EXISTS `snv_pango_marker`;\n"
        "SET FOREIGN_KEY_CHECKS=1;\n",
        encoding="utf-8",
    )


def write_indel_drop_sql() -> None:
    OUT_INDEL_DROP_SQL.write_text(
        "-- Reverse sql/pango_indel_marker.sql without touching snv_pango_marker or mutations.\n"
        "SET FOREIGN_KEY_CHECKS=0;\n"
        "DROP TABLE IF EXISTS `pango_indel_marker`;\n"
        "SET FOREIGN_KEY_CHECKS=1;\n",
        encoding="utf-8",
    )


def _value_tuple(row: dict[str, object], include_kind: bool) -> str:
    parts = [
        sql_str(row["lineage"]),
        sql_str(row["chrom"]),
        str(int(row["coordinate"])),
        sql_str(row["reference"]),
        sql_str(row["alternate"]),
    ]
    if include_kind:
        parts.append(sql_str(row["kind"]))
    parts.extend(
        [
            "NULL" if row["adref"] is None else str(int(row["adref"])),
            "NULL" if row["adalt"] is None else str(int(row["adalt"])),
            "NULL" if row["dp"] is None else str(int(row["dp"])),
            sql_num(row["ratio"]),
            str(int(row["source_row"])),
        ]
    )
    return ",".join(parts)


def write_sql(path: Path, rows: list[dict[str, object]]) -> None:
    lines = [
        "-- Pango lineage names attached to SNVs (Jim Kelley designation markers v1.9).",
        "-- Generated by scripts/import_pango_snv_markers.py",
        "-- Source: " + path.name,
        "-- Rule: len(ref)=1 AND len(alt)=1 AND ratio < 0.2; Column A is the lineage.",
        "-- Reverse with sql/snv_pango_marker_drop.sql",
        "",
        "SET NAMES utf8mb4;",
        "SET FOREIGN_KEY_CHECKS=0;",
        "DROP TABLE IF EXISTS `snv_pango_marker`;",
        "SET FOREIGN_KEY_CHECKS=1;",
        "",
        "CREATE TABLE `snv_pango_marker` (",
        "  `id` int(10) UNSIGNED NOT NULL AUTO_INCREMENT,",
        "  `lineage` varchar(64) NOT NULL,",
        "  `chrom` varchar(32) DEFAULT NULL,",
        "  `coordinate` int(11) NOT NULL,",
        "  `reference` char(1) NOT NULL,",
        "  `alternate` char(1) NOT NULL,",
        "  `adref` int(11) DEFAULT NULL,",
        "  `adalt` int(11) DEFAULT NULL,",
        "  `dp` int(11) DEFAULT NULL,",
        "  `ratio` double DEFAULT NULL,",
        "  `source_row` int(11) NOT NULL,",
        "  PRIMARY KEY (`id`),",
        "  UNIQUE KEY `uq_snv_pango` (`coordinate`,`reference`,`alternate`,`lineage`),",
        "  KEY `idx_snv_pango_snv` (`coordinate`,`reference`,`alternate`)",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;",
        "",
    ]
    if rows:
        lines.append(
            "INSERT INTO `snv_pango_marker` "
            "(`id`,`lineage`,`chrom`,`coordinate`,`reference`,`alternate`,`adref`,`adalt`,`dp`,`ratio`,`source_row`)"
        )
        lines.append("VALUES")
        value_lines = []
        for i, row in enumerate(rows, start=1):
            comma = "," if i < len(rows) else ";"
            value_lines.append("  ({id},{rest}){comma}".format(id=i, rest=_value_tuple(row, False), comma=comma))
        lines.extend(value_lines)
        lines.append("")

    OUT_SQL.write_text("\n".join(lines), encoding="utf-8")


def write_indel_sql(path: Path, rows: list[dict[str, object]]) -> None:
    lines = [
        "-- Pango lineage names attached to indels (Jim Kelley designation markers v1.9).",
        "-- Generated by scripts/import_pango_snv_markers.py",
        "-- Source: " + path.name,
        "-- Rule: len(ref)!=len(alt) AND ratio < 0.2; len(ref)>len(alt) deletion, else insertion.",
        "-- Reverse with sql/pango_indel_marker_drop.sql",
        "",
        "SET NAMES utf8mb4;",
        "SET FOREIGN_KEY_CHECKS=0;",
        "DROP TABLE IF EXISTS `pango_indel_marker`;",
        "SET FOREIGN_KEY_CHECKS=1;",
        "",
        "CREATE TABLE `pango_indel_marker` (",
        "  `id` int(10) UNSIGNED NOT NULL AUTO_INCREMENT,",
        "  `lineage` varchar(64) NOT NULL,",
        "  `chrom` varchar(32) DEFAULT NULL,",
        "  `coordinate` int(11) NOT NULL,",
        "  `reference` varchar(64) NOT NULL,",
        "  `alternate` varchar(64) NOT NULL,",
        "  `kind` enum('deletion','insertion') NOT NULL,",
        "  `adref` int(11) DEFAULT NULL,",
        "  `adalt` int(11) DEFAULT NULL,",
        "  `dp` int(11) DEFAULT NULL,",
        "  `ratio` double DEFAULT NULL,",
        "  `source_row` int(11) NOT NULL,",
        "  PRIMARY KEY (`id`),",
        "  UNIQUE KEY `uq_pango_indel` (`coordinate`,`reference`,`alternate`,`lineage`),",
        "  KEY `idx_pango_indel_coord` (`coordinate`)",
        ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;",
        "",
    ]
    if rows:
        lines.append(
            "INSERT INTO `pango_indel_marker` "
            "(`id`,`lineage`,`chrom`,`coordinate`,`reference`,`alternate`,`kind`,"
            "`adref`,`adalt`,`dp`,`ratio`,`source_row`)"
        )
        lines.append("VALUES")
        value_lines = []
        for i, row in enumerate(rows, start=1):
            comma = "," if i < len(rows) else ";"
            value_lines.append("  ({id},{rest}){comma}".format(id=i, rest=_value_tuple(row, True), comma=comma))
        lines.extend(value_lines)
        lines.append("")

    OUT_INDEL_SQL.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    path = find_xlsx()
    snvs, indels = iter_designation_rows(path)
    snv_keys = {(r["coordinate"], r["reference"], r["alternate"]) for r in snvs}
    indel_keys = {(r["coordinate"], r["reference"], r["alternate"], r["kind"]) for r in indels}
    write_drop_sql()
    write_indel_drop_sql()
    write_sql(path, snvs)
    write_indel_sql(path, indels)
    print("source", path)
    print("qualifying_snv_rows", len(snvs))
    print("distinct_snvs", len(snv_keys))
    print("qualifying_indel_rows", len(indels))
    print("distinct_indels", len(indel_keys))
    print("wrote", OUT_SQL)
    print("wrote", OUT_DROP_SQL)
    print("wrote", OUT_INDEL_SQL)
    print("wrote", OUT_INDEL_DROP_SQL)


if __name__ == "__main__":
    main()
